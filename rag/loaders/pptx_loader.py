from pptx import Presentation


def load_pptx(file_path: str) -> str:
    presentation = Presentation(file_path)

    slides_text = []

    for slide in presentation.slides:
        slide_text = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()

                if text:
                    slide_text.append(text)

        if slide_text:
            slides_text.append("\n".join(slide_text))

    return "\n\n".join(slides_text)